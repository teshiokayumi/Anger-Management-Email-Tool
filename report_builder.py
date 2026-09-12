"""分析結果の可視化（グラフ）と、学年主任との打ち合わせ用資料（Word / Markdown）の作成"""
from collections import Counter
from datetime import date
from io import BytesIO

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from docx import Document
from matplotlib import font_manager, ticker
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from issue_agent import AFTER_HOURS_END, AFTER_HOURS_START, HIGH_STRESS, AgentResult

# ---------- グラフの共通設定 ----------
_JP_FONTS = ["Yu Gothic", "Meiryo", "MS Gothic", "Noto Sans CJK JP", "Noto Sans JP", "IPAexGothic"]
_available = {f.name for f in font_manager.fontManager.ttflist}
plt.rcParams["font.family"] = [f for f in _JP_FONTS if f in _available] or ["sans-serif"]

BLUE = "#2a78d6"
INK = "#0b0b0b"
INK2 = "#52514e"
MUTED = "#898781"
GRID = "#e1e0d9"
AXIS = "#c3c2b7"


def _style(ax, grid_axis="y"):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(AXIS)
    ax.tick_params(colors=INK2, labelsize=9, length=0)
    ax.grid(axis=grid_axis, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def fig_category(cases: list[dict]):
    """分類ごとの件数（横棒）"""
    counts = Counter(c["category"] for c in cases).most_common()[::-1]
    labels, values = zip(*counts)
    fig, ax = plt.subplots(figsize=(6.4, 0.45 * len(labels) + 0.9))
    ax.barh(labels, values, color=BLUE, height=0.6)
    for y, v in enumerate(values):
        ax.text(v + 0.05, y, f"{v}件", va="center", fontsize=9, color=INK2)
    ax.set_xlim(0, max(values) * 1.2)
    ax.xaxis.set_major_locator(ticker.MaxNLocator(integer=True))
    ax.set_title("相談内容の分類", loc="left", fontsize=11, color=INK)
    _style(ax, grid_axis="x")
    fig.tight_layout()
    return fig


def fig_stress(cases: list[dict]):
    """案件ごとの担任のストレス度の推移（折れ線）"""
    xs = [c["created_at"] for c in cases]
    ys = [c["stress"] for c in cases]
    fig, ax = plt.subplots(figsize=(6.4, 3.0))
    ax.plot(xs, ys, color=BLUE, linewidth=2, marker="o", markersize=7,
            markeredgecolor="white", markeredgewidth=1.5)
    ax.axhline(HIGH_STRESS, color=MUTED, linewidth=1, linestyle="--")
    ax.text(1.0, HIGH_STRESS + 2, f"要注意ライン（{HIGH_STRESS}）", transform=ax.get_yaxis_transform(),
            ha="right", va="bottom", fontsize=8, color=INK2)
    for c in cases:
        if c["stress"] >= HIGH_STRESS:
            ax.annotate(c["case_id"], (c["created_at"], c["stress"]), textcoords="offset points",
                        xytext=(0, 8), ha="center", fontsize=8, color=INK2)
    ax.set_ylim(0, 105)
    ax.xaxis.set_major_locator(mdates.AutoDateLocator(minticks=3, maxticks=6))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%d"))
    ax.set_title("担任のストレス度の推移（0〜100）", loc="left", fontsize=11, color=INK)
    _style(ax)
    fig.tight_layout()
    return fig


def fig_issue_matrix(result: AgentResult):
    """課題の優先度マップ（緊急度 × 影響度）"""
    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    ax.axvspan(3, 5.5, ymin=0.5, ymax=1, color="#f3f7fd", zorder=0)
    ax.axvline(3, color=AXIS, linewidth=1, linestyle="--")
    ax.axhline(3, color=AXIS, linewidth=1, linestyle="--")
    ax.text(5.4, 5.4, "優先して取り組む", ha="right", va="top", fontsize=9, color=INK2)

    placed = Counter()
    for i, issue in enumerate(result.analysis.issues, 1):
        key = (issue.urgency, issue.impact)
        offset = placed[key] * 0.45  # 同じ位置に重なる課題は下にずらす
        placed[key] += 1
        x, y = issue.urgency, issue.impact - offset
        ax.scatter(x, y, s=360, color=BLUE, edgecolors="white", linewidths=2, zorder=3)
        ax.text(x, y, str(i), ha="center", va="center", fontsize=10, color="white", weight="bold", zorder=4)
        title = issue.title if len(issue.title) <= 14 else issue.title[:13] + "…"
        # 右端の課題はラベルを左側に置いて、はみ出さないようにする
        right = x < 4
        ax.text(x + (0.22 if right else -0.22), y, title, ha="left" if right else "right", va="center",
                fontsize=8.5, color=INK, zorder=4)

    ax.set_xlim(0.5, 5.5)
    ax.set_ylim(0.5, 5.5)
    ax.set_xticks(range(1, 6))
    ax.set_yticks(range(1, 6))
    ax.set_xlabel("緊急度 →", fontsize=9, color=INK2)
    ax.set_ylabel("影響の大きさ →", fontsize=9, color=INK2)
    ax.set_title("課題の優先度マップ", loc="left", fontsize=11, color=INK)
    _style(ax, grid_axis="both")
    fig.tight_layout()
    return fig


def fig_to_png(fig) -> bytes:
    buf = BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ---------- 資料の共通データ ----------
def metric_rows(result: AgentResult) -> list[tuple[str, str]]:
    s = result.stats
    repeat = "、".join(f"{k}（{v}件）" for k, v in s["repeat_parents"].items()) or "なし"
    return [
        ("保護者対応の件数", f"{s['total']}件"),
        ("担任の平均ストレス度", f"{s['avg_stress']} / 100"),
        (f"ストレス度{HIGH_STRESS}以上の案件", f"{s['high_stress']}件"),
        ("学年・管理職と共有すべき案件", f"{s['escalation']}件"),
        (f"{AFTER_HOURS_START}時以降・{AFTER_HOURS_END}時前に作成した下書き", f"{s['after_hours']}件"),
        ("複数回やりとりのある保護者", repeat),
    ]


def issue_title(result: AgentResult, issue_no: int) -> str:
    issues = result.analysis.issues
    return f"{issue_no}. {issues[issue_no - 1].title}" if 1 <= issue_no <= len(issues) else "―"


def chart_pngs(result: AgentResult) -> dict[str, bytes]:
    """画面表示とWord資料で共通に使うグラフ画像"""
    return {
        "category": fig_to_png(fig_category(result.cases)),
        "stress": fig_to_png(fig_stress(result.cases)),
        "matrix": fig_to_png(fig_issue_matrix(result)),
    }


# ---------- Word ----------
def _jp_font(style, size: float | None = None, bold: bool | None = None):
    style.font.name = "Yu Gothic"
    style.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rfonts = style.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:eastAsia"), "游ゴシック")


def _shade(cell, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table(doc, header: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text, width in zip(table.rows[0].cells, header, widths_cm):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        cell.width = Cm(width)
        _shade(cell, "EEF3FA")
    for row in rows:
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths_cm):
            cell.text = str(text)
            cell.width = Cm(width)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return table


def build_docx(result: AgentResult, charts: dict[str, bytes], summary: str, approaches: list[dict],
               consultation: list[str], author: str = "", grade: str = "") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)

    _jp_font(doc.styles["Normal"], 10.5)
    _jp_font(doc.styles["Title"], 18, True)
    _jp_font(doc.styles["Heading 1"], 13, True)
    _jp_font(doc.styles["Heading 2"], 11, True)

    header = section.header.paragraphs[0]
    header.text = "取扱注意（校内限り）"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    doc.add_paragraph("保護者対応に関する相談資料", style="Title")
    start, end = result.period
    info = [f"作成日：{date.today():%Y年%m月%d日}", f"対象期間：{start:%m/%d}〜{end:%m/%d}"]
    if grade:
        info.append(f"学年・学級：{grade}")
    if author:
        info.append(f"作成者：{author}")
    doc.add_paragraph("　／　".join(info))

    doc.add_heading("1. 概要", level=1)
    doc.add_paragraph(summary)
    _table(doc, ["指標", "値"], [list(r) for r in metric_rows(result)], [8.5, 8.5])

    doc.add_heading("2. 状況の可視化", level=1)
    for png in charts.values():
        doc.add_picture(BytesIO(png), width=Cm(15))

    doc.add_heading("3. 見えてきた課題", level=1)
    for i, issue in enumerate(result.analysis.issues, 1):
        doc.add_heading(f"課題{i}：{issue.title}", level=2)
        doc.add_paragraph(issue.description)
        doc.add_paragraph(f"考えられる原因：{issue.root_cause}")
        doc.add_paragraph(f"緊急度 {issue.urgency} ／ 影響 {issue.impact}　根拠：{'、'.join(issue.evidence_case_ids)}")

    doc.add_heading("4. 対応アプローチ案", level=1)
    rows = [[issue_title(result, int(a["issue_no"])), a["action"], a["owner"], a["timeframe"], a["expected_effect"]]
            for a in approaches]
    _table(doc, ["課題", "取り組み", "担当", "時期", "期待される効果"], rows, [3.5, 5.5, 2.2, 2.2, 3.6])

    doc.add_heading("5. 学年主任に相談したいこと", level=1)
    for point in consultation:
        doc.add_paragraph(point, style="List Number")

    doc.add_heading("6. 担任の負担について", level=1)
    doc.add_paragraph(result.analysis.stress_trend)
    doc.add_paragraph(result.plan.teacher_care)

    doc.add_heading("付録：案件一覧", level=1)
    rows = [[c["case_id"], f"{c['created_at']:%m/%d %H:%M}", c["category"], c["parent_label"], c["summary"],
             str(c["urgency"]), "要" if c["needs_escalation"] else "―"] for c in result.cases]
    _table(doc, ["ID", "日時", "分類", "相手", "概要", "緊急度", "共有"], rows, [1.2, 2.2, 2.6, 1.8, 6.8, 1.2, 1.2])

    note = doc.add_paragraph("※ 本資料は、メール下書きをAIが分析した結果を担任が確認・修正したものです。"
                             "児童生徒・保護者の氏名は匿名化しています。")
    note.runs[0].font.size = Pt(8.5)
    note.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Markdown ----------
def build_markdown(result: AgentResult, summary: str, approaches: list[dict], consultation: list[str],
                   author: str = "", grade: str = "") -> str:
    start, end = result.period
    lines = ["# 保護者対応に関する相談資料", "",
             f"作成日：{date.today():%Y年%m月%d日}　対象期間：{start:%m/%d}〜{end:%m/%d}"
             + (f"　学年・学級：{grade}" if grade else "") + (f"　作成者：{author}" if author else ""),
             "", "## 1. 概要", "", summary, "", "| 指標 | 値 |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in metric_rows(result)]
    lines += ["", "## 2. 見えてきた課題", ""]
    for i, issue in enumerate(result.analysis.issues, 1):
        lines += [f"### 課題{i}：{issue.title}", "", issue.description, "",
                  f"- 考えられる原因：{issue.root_cause}",
                  f"- 緊急度 {issue.urgency} ／ 影響 {issue.impact}　根拠：{'、'.join(issue.evidence_case_ids)}", ""]
    lines += ["## 3. 対応アプローチ案", "", "| 課題 | 取り組み | 担当 | 時期 | 期待される効果 |", "|---|---|---|---|---|"]
    lines += [f"| {issue_title(result, int(a['issue_no']))} | {a['action']} | {a['owner']} | {a['timeframe']} | "
              f"{a['expected_effect']} |" for a in approaches]
    lines += ["", "## 4. 学年主任に相談したいこと", ""]
    lines += [f"{i}. {p}" for i, p in enumerate(consultation, 1)]
    lines += ["", "## 5. 担任の負担について", "", result.analysis.stress_trend, "", result.plan.teacher_care, "",
              "---", "※ AIの分析結果を担任が確認・修正したものです。氏名は匿名化しています。"]
    return "\n".join(lines)
